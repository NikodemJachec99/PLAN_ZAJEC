"""Read-only view over an openpyxl worksheet.

Merged ranges are expanded so every covered cell reports the value *and* style of the
range's top-left cell. The practical-classes matrix relies on this heavily: a cell merged
across subgroup rows "a" and "b" means both subgroups have that class, and losing the
second row is exactly how entries used to disappear from the plan.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any, Iterator

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from .textutil import clean


@dataclass(frozen=True)
class CellStyle:
    underline: str | None
    strike: bool
    italic: bool
    fill: str | None

    @property
    def time_key(self) -> tuple[str, bool, bool]:
        """Font features planners use to encode working hours (e.g. double underline)."""
        return (self.underline or "none", self.strike, self.italic)


EMPTY_STYLE = CellStyle(underline=None, strike=False, italic=False, fill=None)


def _color_key(color: Any) -> str | None:
    if color is None:
        return None
    kind = getattr(color, "type", None)
    if kind == "rgb":
        rgb = getattr(color, "rgb", None)
        if isinstance(rgb, str) and len(rgb) >= 6:
            return rgb[-6:].upper()
        return None
    if kind == "theme":
        tint = getattr(color, "tint", 0.0) or 0.0
        return f"theme{color.theme}:{round(float(tint), 2)}"
    if kind == "indexed":
        return f"idx{color.indexed}"
    return None


def _style_of(cell: Any) -> CellStyle:
    font = cell.font
    fill = cell.fill
    fill_key = None
    if fill is not None and fill.fill_type:
        fill_key = _color_key(fill.fgColor)
        if fill_key in {"000000", "idx64"} and fill.fill_type == "solid":
            # "Automatic" foreground: fall back to the background colour.
            fill_key = _color_key(fill.bgColor) or fill_key
    underline = font.underline if font is not None else None
    return CellStyle(
        underline=str(underline).lower() if underline else None,
        strike=bool(font.strike) if font is not None else False,
        italic=bool(font.i) if font is not None else False,
        fill=fill_key,
    )


class SheetGrid:
    def __init__(self, worksheet: Worksheet) -> None:
        self.title = worksheet.title
        self.hidden = worksheet.sheet_state != "visible"
        self.max_row = worksheet.max_row or 0
        self.max_col = worksheet.max_column or 0
        self._ws = worksheet
        self._values: dict[tuple[int, int], Any] = {}
        self._origin: dict[tuple[int, int], tuple[int, int]] = {}
        self._style_cache: dict[tuple[int, int], CellStyle] = {}

        for row in worksheet.iter_rows(min_row=1, max_row=self.max_row, max_col=self.max_col):
            for cell in row:
                if cell.value is not None:
                    self._values[(cell.row, cell.column)] = cell.value

        for merged in worksheet.merged_cells.ranges:
            top, left = merged.min_row, merged.min_col
            value = self._values.get((top, left))
            for r in range(merged.min_row, merged.max_row + 1):
                for c in range(merged.min_col, merged.max_col + 1):
                    self._origin[(r, c)] = (top, left)
                    if value is not None:
                        self._values[(r, c)] = value
                    else:
                        self._values.pop((r, c), None)

    def value(self, row: int, col: int) -> Any:
        return self._values.get((row, col))

    def text(self, row: int, col: int) -> str:
        return clean(self._values.get((row, col)))

    def origin(self, row: int, col: int) -> tuple[int, int]:
        return self._origin.get((row, col), (row, col))

    def style(self, row: int, col: int) -> CellStyle:
        key = self.origin(row, col)
        cached = self._style_cache.get(key)
        if cached is None:
            try:
                cached = _style_of(self._ws.cell(row=key[0], column=key[1]))
            except Exception:  # pragma: no cover - defensive: odd styles must not kill parsing
                cached = EMPTY_STYLE
            self._style_cache[key] = cached
        return cached

    def row_texts(self, row: int, *, origins_only: bool = False) -> Iterator[tuple[int, str]]:
        for col in range(1, self.max_col + 1):
            if origins_only and self.origin(row, col) != (row, col):
                continue
            text = self.text(row, col)
            if text:
                yield col, text

    def is_row_empty(self, row: int) -> bool:
        return not any(True for _ in self.row_texts(row))


def load_grids(source: Path | bytes) -> list[SheetGrid]:
    stream: Any = BytesIO(source) if isinstance(source, (bytes, bytearray)) else source
    workbook = openpyxl.load_workbook(stream, data_only=True)
    return [SheetGrid(sheet) for sheet in workbook.worksheets]
